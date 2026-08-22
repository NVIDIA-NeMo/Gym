# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
# http://www.apache.org/licenses/LICENSE-2.0
"""Pre-convert Office documents to PDF for GDPVal judging.

Each invocation gets its own ``-env:UserInstallation`` profile dir, so
concurrent libreoffice subprocesses don't race on the shared default
profile lock (``$HOME/.config/libreoffice``) — that race is the reason
the previous default ``max_concurrent=1`` existed.

OOXML namespace normalization
-----------------------------

Some files in the GDPVal corpus were emitted by ``python-docx`` (or
similar lxml-based tools), which serialize the OPC package XML with an
explicit ``ns0:`` namespace prefix:

    <ns0:Relationships xmlns:ns0="http://schemas.openxmlformats.org/...">

instead of the standard default-namespace form:

    <Relationships xmlns="http://schemas.openxmlformats.org/...">

The two forms are semantically identical XML, and Microsoft Word /
pandoc accept both. LibreOffice 24.2, however, rejects the prefixed
form with ``Error: source file could not be loaded``. The prefixing
shows up in BOTH ``_rels/.rels`` and ``[Content_Types].xml``; rewriting
only one of them is not enough.

Before invoking libreoffice we detect this shape and write a
namespace-normalized copy to a tempdir, leaving the original on disk
untouched.
"""

from __future__ import annotations

import asyncio
import logging
import os
import re
import shutil
import subprocess
import tempfile
import zipfile
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable


LOGGER = logging.getLogger(__name__)

OOXML_EXTENSIONS = {".docx", ".pptx", ".xlsx"}
# LibreOffice converts these fine; omitting them left the judge with no PDF.
# The ns0 pre-pass skips them harmlessly (OLE files raise BadZipFile).
LEGACY_OFFICE_EXTENSIONS = {".doc", ".ppt", ".xls"}

OFFICE_EXTENSIONS = OOXML_EXTENSIONS | LEGACY_OFFICE_EXTENSIONS

DEFAULT_MAX_CONCURRENT = 4

_NS0_ROOT_RE = re.compile(r'<ns0:([A-Za-z_][\w.-]*)\b([^>]*?)\bxmlns:ns0="([^"]+)"')
_NS0_TAG_RE = re.compile(r"</?ns0:")
_NS0_SENTINEL = b'xmlns:ns0="http://schemas.openxmlformats.org/'


def _rewrite_ns0_namespace(text: str) -> str:
    text = _NS0_ROOT_RE.sub(r'<\1 xmlns="\3"\2', text)
    text = _NS0_TAG_RE.sub(lambda m: m.group(0).replace("ns0:", ""), text)
    return text


def _ooxml_has_ns0_prefix(path: Path) -> bool:
    """True if the package uses python-docx-style ``ns0:`` prefixing in
    ``_rels/.rels`` or ``[Content_Types].xml``. LibreOffice can't load
    files in this form even though they are valid OOXML."""
    try:
        with zipfile.ZipFile(path) as zin:
            names = set(zin.namelist())
            for part in ("_rels/.rels", "[Content_Types].xml"):
                if part in names and _NS0_SENTINEL in zin.read(part):
                    return True
    except (zipfile.BadZipFile, OSError):
        return False
    return False


def _normalize_ooxml_zip(src: Path, dst: Path) -> None:
    """Copy ``src`` to ``dst`` rewriting any ``ns0:``-prefixed package XML
    (``*.rels`` and ``[Content_Types].xml``) to default-namespace form."""
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.namelist():
            data = zin.read(item)
            if item.endswith(".rels") or item == "[Content_Types].xml":
                data = _rewrite_ns0_namespace(data.decode("utf-8")).encode("utf-8")
            zout.writestr(item, data)


def sidecar_pdf(path: Path) -> Path:
    """Injective PDF name for *path*: ``Plan.pptx`` -> ``Plan.pptx.pdf``."""
    return path.with_name(path.name + ".pdf")


@dataclass(frozen=True)
class PdfProvenance:
    """Existing Office renders and PDFs that must not be emitted standalone.

    A plain ``Plan.pdf`` cannot identify which source it renders when both
    ``Plan.docx`` and ``Plan.pptx`` exist.  In that case it is quarantined as
    ambiguous and only injective ``Plan.docx.pdf`` / ``Plan.pptx.pdf`` sidecars
    are trusted.  PDFs whose stem does not collide with an Office source are
    never included in :attr:`suppressed_pdfs` and remain independent artifacts.
    """

    office_pdfs: dict[Path, Path]
    suppressed_pdfs: frozenset[Path]
    ambiguous_pdfs: frozenset[Path]


def base64_encoded_size(raw_bytes: int) -> int:
    """Exact base64 payload length for *raw_bytes* without encoding the data."""

    return 4 * ((max(0, raw_bytes) + 2) // 3)


@dataclass
class AttachmentBudget:
    """Mutable raw/encoded byte budget used before loading judge attachments."""

    raw_limit: int
    encoded_limit: int
    raw_used: int = 0
    encoded_used: int = 0

    def can_fit(self, raw_bytes: int, encoded_chars: int | None = None) -> bool:
        raw = max(0, raw_bytes)
        encoded = base64_encoded_size(raw) if encoded_chars is None else max(0, encoded_chars)
        return self.raw_used + raw <= self.raw_limit and self.encoded_used + encoded <= self.encoded_limit

    def reserve(self, raw_bytes: int, encoded_chars: int | None = None) -> bool:
        raw = max(0, raw_bytes)
        encoded = base64_encoded_size(raw) if encoded_chars is None else max(0, encoded_chars)
        if not self.can_fit(raw, encoded):
            return False
        self.raw_used += raw
        self.encoded_used += encoded
        return True


def resolve_pdf_provenance(paths: Iterable[Path]) -> PdfProvenance:
    """Resolve sidecar PDFs for a directory listing without guessing provenance.

    The injective ``source.ext.pdf`` spelling always wins.  The historical
    ``source.pdf`` spelling is accepted only when exactly one Office source in
    the same directory has that stem.  Any PDF selected as an Office render is
    consumed, and an ambiguous historical sibling is suppressed so stale cache
    artifacts cannot be judged a second (or third) time.
    """

    entries = [Path(path) for path in paths]
    entry_set = set(entries)
    groups: dict[tuple[Path, str], list[Path]] = {}
    for path in entries:
        if path.suffix.lower() in OFFICE_EXTENSIONS:
            groups.setdefault((path.parent, path.stem), []).append(path)

    office_pdfs: dict[Path, Path] = {}
    suppressed: set[Path] = set()
    ambiguous: set[Path] = set()
    for (parent, stem), sources in groups.items():
        plain_pdf = parent / f"{stem}.pdf"
        plain_exists = plain_pdf in entry_set and plain_pdf.is_file()
        is_ambiguous = len(sources) > 1

        if plain_exists:
            suppressed.add(plain_pdf)
            if is_ambiguous:
                ambiguous.add(plain_pdf)

        for source in sources:
            injective = sidecar_pdf(source)
            if injective in entry_set and injective.is_file():
                office_pdfs[source] = injective
                suppressed.add(injective)
            elif plain_exists and not is_ambiguous:
                office_pdfs[source] = plain_pdf

    return PdfProvenance(
        office_pdfs=office_pdfs,
        suppressed_pdfs=frozenset(suppressed),
        ambiguous_pdfs=frozenset(ambiguous),
    )


def roundtrip_ooxml_copy(src: Path, dst: Path) -> None:
    """Open and re-save an OOXML package to *dst*, leaving *src* unchanged.

    This is deliberately a retry repair rather than a pre-pass: most Office
    documents convert directly, while a small class of valid OPC packages is
    rejected by LibreOffice until the format-specific library rewrites it.
    """

    extension = src.suffix.lower()
    if extension == ".docx":
        from docx import Document

        Document(str(src)).save(str(dst))
    elif extension == ".pptx":
        from pptx import Presentation

        Presentation(str(src)).save(str(dst))
    elif extension == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(str(src))
        try:
            workbook.save(str(dst))
        finally:
            workbook.close()
    else:
        raise ValueError(f"OOXML round-trip is unsupported for {src.suffix or '<no extension>'}")


def extract_xlsx_structured_text(path: Path, *, max_chars: int = 120_000) -> str:
    """Return bounded cell-level XLSX evidence including formulas and values.

    Sheet XML is streamed directly rather than iterating openpyxl's rectangular
    cell range. That matters for sparse workbooks whose used range says e.g.
    ``XFD1048576``: visiting every empty coordinate would take effectively
    forever. OOXML stores formulas and cached values together, so this also
    preserves both when the producer wrote a cache.
    """

    import posixpath
    import xml.etree.ElementTree as element_tree

    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    document_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    cell_tag = f"{{{spreadsheet_ns}}}c"
    row_tag = f"{{{spreadsheet_ns}}}row"
    formula_tag = f"{{{spreadsheet_ns}}}f"
    value_tag = f"{{{spreadsheet_ns}}}v"
    inline_string_tag = f"{{{spreadsheet_ns}}}is"
    text_tag = f"{{{spreadsheet_ns}}}t"

    # Records are bounded before shared-string resolution, so even a workbook
    # with millions of populated cells cannot grow this function's output-side
    # memory without limit.
    records: list[tuple[str, str, str | None, str | None, str | None]] = []
    shared_indices: set[int] = set()
    estimated_chars = 0
    truncated = False

    def _record_cell(sheet_name: str, cell: element_tree.Element) -> None:
        nonlocal estimated_chars, truncated
        coordinate = cell.attrib.get("r", "?")
        cell_type = cell.attrib.get("t", "")
        formula_node = cell.find(formula_tag)
        value_node = cell.find(value_tag)
        inline_node = cell.find(inline_string_tag)
        formula = formula_node.text if formula_node is not None else None
        value = value_node.text if value_node is not None else None
        inline = "".join(node.text or "" for node in inline_node.iter(text_tag)) if inline_node is not None else None
        if formula is None and value is None and not inline:
            return

        # Cap stored raw strings too; a single cell can contain far more text
        # than the whole judge budget.
        room = max(0, max_chars - estimated_chars)
        formula = formula[:room] if formula is not None else None
        value = value[:room] if value is not None else None
        inline = inline[:room] if inline is not None else None
        records.append((sheet_name, coordinate, cell_type, formula, inline if inline is not None else value))
        estimated_chars += min(
            room,
            len(coordinate) + len(formula or "") + len(inline or value or "") + 20,
        )
        if cell_type == "s" and value is not None:
            try:
                shared_indices.add(int(value))
            except ValueError:
                pass
        if estimated_chars >= max_chars:
            truncated = True

    with zipfile.ZipFile(path) as archive:
        workbook = element_tree.fromstring(archive.read("xl/workbook.xml"))
        relationships = element_tree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        targets = {
            relation.attrib["Id"]: relation.attrib["Target"]
            for relation in relationships.findall(f"{{{package_rel_ns}}}Relationship")
        }

        sheets: list[tuple[str, str]] = []
        for sheet in workbook.findall(f".//{{{spreadsheet_ns}}}sheet"):
            relationship_id = sheet.attrib.get(f"{{{document_rel_ns}}}id")
            target = targets.get(relationship_id or "")
            if not target:
                continue
            part = target.lstrip("/") if target.startswith("/") else posixpath.normpath(f"xl/{target}")
            sheets.append((sheet.attrib.get("name", "unnamed"), part))

        for sheet_name, part in sheets:
            header_cost = len(sheet_name) + 9
            if estimated_chars + header_cost > max_chars:
                truncated = True
                break
            estimated_chars += header_cost
            try:
                stream = archive.open(part)
            except KeyError:
                continue
            with stream:
                stack: list[element_tree.Element] = []
                for event, node in element_tree.iterparse(stream, events=("start", "end")):
                    if event == "start":
                        stack.append(node)
                        continue
                    parent = stack[-2] if len(stack) > 1 else None
                    if node.tag == cell_tag:
                        _record_cell(sheet_name, node)
                        node.clear()
                    elif node.tag == row_tag:
                        # Clearing cells alone leaves one empty element per cell
                        # attached to its row. Remove the completed row from
                        # sheetData as well so memory stays bounded by one row.
                        node.clear()
                        if parent is not None:
                            parent.remove(node)
                    stack.pop()
                    if truncated:
                        break
            if truncated:
                break

        shared_strings: dict[int, str] = {}
        if shared_indices and "xl/sharedStrings.xml" in archive.namelist():
            wanted = set(shared_indices)
            shared_index = -1
            shared_chars_remaining = max_chars
            with archive.open("xl/sharedStrings.xml") as stream:
                stack = []
                for event, node in element_tree.iterparse(stream, events=("start", "end")):
                    if event == "start":
                        stack.append(node)
                        continue
                    parent = stack[-2] if len(stack) > 1 else None
                    done = False
                    if node.tag == f"{{{spreadsheet_ns}}}si":
                        shared_index += 1
                        if shared_index in wanted:
                            pieces: list[str] = []
                            remaining = shared_chars_remaining
                            for text_node in node.iter(text_tag):
                                piece = text_node.text or ""
                                pieces.append(piece[:remaining])
                                remaining -= min(len(piece), remaining)
                                if remaining <= 0:
                                    break
                            shared_strings[shared_index] = "".join(pieces)
                            shared_chars_remaining -= len(shared_strings[shared_index])
                            wanted.remove(shared_index)
                            done = not wanted or shared_chars_remaining <= 0
                        node.clear()
                        if parent is not None:
                            parent.remove(node)
                    stack.pop()
                    if done:
                        break

    lines: list[str] = []
    used = 0
    current_sheet: str | None = None

    def _append(line: str, *, blank_before: bool = False) -> bool:
        nonlocal used, truncated
        separator = 2 if lines and blank_before else (1 if lines else 0)
        if used + separator + len(line) > max_chars:
            truncated = True
            available = max(0, max_chars - used - separator)
            if available > 0:
                if lines and blank_before:
                    lines.append("")
                lines.append(line[:available])
                used += separator + available
            return False
        if lines and blank_before:
            lines.append("")
        lines.append(line)
        used += separator + len(line)
        return True

    for sheet_name, coordinate, cell_type, formula, raw_value in records:
        if sheet_name != current_sheet:
            if not _append(f"Sheet: {sheet_name}", blank_before=True):
                break
            current_sheet = sheet_name

        value = raw_value
        if cell_type == "s" and raw_value is not None:
            try:
                value = shared_strings.get(int(raw_value), f"[shared string #{raw_value}]")
            except ValueError:
                pass
        elif cell_type == "b" and raw_value is not None:
            value = "TRUE" if raw_value == "1" else "FALSE"

        if formula is not None:
            rendered = f"{coordinate}: formula: ={formula.lstrip('=')}"
            if value not in (None, ""):
                rendered += f"; cached/display value: {value}"
        else:
            rendered = f"{coordinate}: value={value}"
        if not _append(rendered):
            break

    marker = "[...spreadsheet text truncated]"
    text = "\n".join(lines)
    if truncated:
        keep = max(0, max_chars - len(marker) - 1)
        text = text[:keep].rstrip() + "\n" + marker
    return text[:max_chars]


def needs_conversion(path: Path, *, ambiguous: bool = False) -> bool:
    if path.suffix.lower() not in OFFICE_EXTENSIONS:
        return False
    if sidecar_pdf(path).exists():
        return False
    if ambiguous:
        return True
    return not path.with_suffix(".pdf").exists()


def _safe_basename(name: str) -> str:
    """Sanitize a filename's stem for LibreOffice's batch-convert mode.

    Whitespace in the input path makes LibreOffice's internal URI handling
    drop arguments mid-flight, producing ``source file could not be loaded``
    with rc=0. Replacing whitespace in the stem (extension preserved) and
    staging via tempdir sidesteps the bug.
    """
    p = Path(name)
    return re.sub(r"\s+", "_", p.stem) + p.suffix


def convert_to_pdf(path: Path, output_pdf: Path | None = None) -> tuple[Path, bool, str]:
    """Convert one file to PDF via host LibreOffice. Returns ``(path, ok, msg)``.

    *output_pdf* overrides the destination: LibreOffice names output after the
    input stem, so same-stem files would otherwise race for one name.
    """
    profile_dir = Path(tempfile.mkdtemp(prefix="lo-profile-"))
    stage_dir: Path | None = None
    retry_dir: Path | None = None
    retry_profile_dir: Path | None = None
    input_path = path
    normalized = False
    needs_ns0_normalization = _ooxml_has_ns0_prefix(path)
    has_whitespace = any(c.isspace() for c in path.name)
    # A custom destination always stages, else both sides write one name.
    needs_stage = needs_ns0_normalization or has_whitespace or output_pdf is not None
    try:
        if needs_stage:
            stage_dir = Path(tempfile.mkdtemp(prefix="gdpval-stage-"))
            staged_name = _safe_basename(path.name) if has_whitespace else path.name
            input_path = stage_dir / staged_name
            if needs_ns0_normalization:
                _normalize_ooxml_zip(path, input_path)
                normalized = True
            else:
                shutil.copy2(path, input_path)
            lo_outdir = str(stage_dir)
        else:
            lo_outdir = str(path.parent)

        def _run_libreoffice(source: Path, outdir: str, profile: Path) -> subprocess.CompletedProcess[str]:
            return subprocess.run(
                [
                    "libreoffice",
                    "--headless",
                    "--nologo",
                    "--nolockcheck",
                    "--nodefault",
                    "--norestore",
                    f"-env:UserInstallation=file://{profile.as_posix()}",
                    "--convert-to",
                    "pdf",
                    "--outdir",
                    outdir,
                    str(source),
                ],
                capture_output=True,
                text=True,
                timeout=120,
            )

        result = _run_libreoffice(input_path, lo_outdir, profile_dir)
        final_pdf = output_pdf if output_pdf is not None else path.with_suffix(".pdf")
        if needs_stage:
            staged_pdf = Path(lo_outdir) / (input_path.stem + ".pdf")
            if staged_pdf.exists():
                shutil.move(str(staged_pdf), str(final_pdf))
        if final_pdf.exists():
            suffix = " (after ns0 normalization)" if normalized else ""
            return path, True, f"converted {path.name}{suffix}"

        # LibreOffice sometimes returns rc=0 while rejecting a valid OOXML OPC
        # package. Re-saving through its format library rewrites the package in
        # a canonical shape. Retry exactly once from that temp copy; the source
        # deliverable is never opened for writing.
        retry_error: str | None = None
        if path.suffix.lower() in OOXML_EXTENSIONS:
            retry_dir = Path(tempfile.mkdtemp(prefix="gdpval-roundtrip-"))
            retry_profile_dir = Path(tempfile.mkdtemp(prefix="lo-profile-retry-"))
            retry_input = retry_dir / _safe_basename(path.name)
            try:
                roundtrip_ooxml_copy(path, retry_input)
                retry_result = _run_libreoffice(retry_input, str(retry_dir), retry_profile_dir)
                retry_pdf = retry_dir / (retry_input.stem + ".pdf")
                if retry_pdf.exists():
                    shutil.move(str(retry_pdf), str(final_pdf))
                    return path, True, f"converted {path.name} (after OOXML round-trip retry)"
                retry_error = f"round-trip retry rc={retry_result.returncode}: {retry_result.stderr.strip()[:300]}"
            except Exception as exc:
                retry_error = f"round-trip retry failed: {exc!r}"

        retry_suffix = f"; {retry_error}" if retry_error else ""
        return (
            path,
            False,
            f"libreoffice rc={result.returncode} did not produce {final_pdf.name}: "
            f"{result.stderr.strip()[:300]}{retry_suffix}",
        )
    except subprocess.TimeoutExpired:
        return path, False, f"timeout converting {path.name}"
    except FileNotFoundError:
        return path, False, "libreoffice not found on host PATH (install with: apt install libreoffice)"
    except Exception as exc:
        return path, False, f"error converting {path.name}: {exc!r}"
    finally:
        shutil.rmtree(profile_dir, ignore_errors=True)
        if stage_dir is not None:
            shutil.rmtree(stage_dir, ignore_errors=True)
        if retry_dir is not None:
            shutil.rmtree(retry_dir, ignore_errors=True)
        if retry_profile_dir is not None:
            shutil.rmtree(retry_profile_dir, ignore_errors=True)


def find_convertible_files(root_dir: str | os.PathLike) -> list[tuple[Path, Path | None]]:
    """Office files needing conversion, as ``(source, explicit_destination)``.

    ``Report.docx`` and ``Report.pptx`` both target ``Report.pdf`` and race for
    it, leaving a PDF that looks like a valid render of both. Same-stem files
    get the injective sidecar name instead; ``file_reader`` prefers it.
    """
    files: list[tuple[Path, Path | None]] = []
    for dirpath, _, filenames in os.walk(root_dir):
        office_stems: dict[str, int] = {}
        for filename in filenames:
            path = Path(dirpath) / filename
            if path.suffix.lower() in OFFICE_EXTENSIONS:
                office_stems[path.stem] = office_stems.get(path.stem, 0) + 1

        for filename in filenames:
            path = Path(dirpath) / filename
            ambiguous = office_stems.get(path.stem, 0) > 1
            if not needs_conversion(path, ambiguous=ambiguous):
                continue
            if ambiguous:
                LOGGER.info(
                    "Preconverting %s to sidecar '%s': another Office file shares its stem",
                    path,
                    sidecar_pdf(path).name,
                )
                files.append((path, sidecar_pdf(path)))
            else:
                files.append((path, None))
    return sorted(files, key=lambda pair: pair[0])


def preconvert_dir(
    root_dir: str | os.PathLike,
    max_concurrent: int = DEFAULT_MAX_CONCURRENT,
) -> tuple[int, int, list[str]]:
    """Convert every pending Office file under ``root_dir`` to PDF.

    Returns ``(num_success, num_failed, error_messages)``. Caller should log
    a sample at WARNING when ``num_failed > 0``.
    """
    files = find_convertible_files(root_dir)
    if not files:
        return 0, 0, []

    success_count = 0
    fail_count = 0
    error_messages: list[str] = []

    # Same-stem files convert SEQUENTIALLY: launched together LibreOffice aborts
    # one side with rc=134, and per-conversion UserInstallation profiles do not
    # prevent it. There are a handful per tree, so this costs almost nothing.
    parallel = [(src, dest) for src, dest in files if dest is None]
    serial = [(src, dest) for src, dest in files if dest is not None]

    def _tally(result: tuple[Path, bool, str]) -> None:
        nonlocal success_count, fail_count
        _, success, message = result
        if success:
            success_count += 1
        else:
            fail_count += 1
            error_messages.append(message)

    if parallel:
        with ThreadPoolExecutor(max_workers=max_concurrent) as executor:
            futures = [executor.submit(convert_to_pdf, src, dest) for src, dest in parallel]
            for future in as_completed(futures):
                _tally(future.result())

    for src, dest in serial:
        _tally(convert_to_pdf(src, dest))

    return success_count, fail_count, error_messages


async def preconvert_dir_async(
    root_dir: str | os.PathLike,
    max_concurrent: int = DEFAULT_MAX_CONCURRENT,
) -> tuple[int, int, list[str]]:
    return await asyncio.to_thread(preconvert_dir, root_dir, max_concurrent)
