# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
"""Provision and verify document rendering before accepting GDPVal tasks."""

import fcntl
import shutil
import subprocess
import tempfile
from pathlib import Path


def ensure_judge_runtime():
    from resources_servers.gdpval.setup_libreoffice import ensure_libreoffice

    # Train/validation resource processes share one container's package database.
    with open(Path(tempfile.gettempdir()) / "gdpval-judge-runtime.lock", "a") as lock:
        fcntl.flock(lock, fcntl.LOCK_EX)
        ready = shutil.which("libreoffice") and shutil.which("java")
        if ready:
            ready = (
                subprocess.run(
                    ["java", "-version"], capture_output=True, timeout=30, check=False
                ).returncode
                == 0
            )
        if not ready and not ensure_libreoffice():
            raise RuntimeError(
                "GDPVal judge requires LibreOffice and a working Java runtime"
            )
        _check_office_rendering()


def _check_office_rendering():
    import fitz
    from docx import Document
    from openpyxl import Workbook

    with tempfile.TemporaryDirectory(prefix="gdpval-judge-render-") as temporary:
        root = Path(temporary)
        document = Document()
        document.add_paragraph("GDPVal judge rendering check")
        document.save(root / "document.docx")
        workbook = Workbook()
        workbook.active.append(["GDPVal", 2])
        workbook.active.append(["Total", "=B1*2"])
        workbook.save(root / "spreadsheet.xlsx")
        result = subprocess.run(
            [
                "libreoffice",
                "--headless",
                "--nologo",
                "--nolockcheck",
                f"-env:UserInstallation={(root / 'profile').as_uri()}",
                "--convert-to",
                "pdf",
                "--outdir",
                str(root),
                str(root / "document.docx"),
                str(root / "spreadsheet.xlsx"),
            ],
            check=False,
            capture_output=True,
            text=True,
            timeout=120,
        )
        if result.returncode:
            raise RuntimeError(
                f"GDPVal judge Office conversion failed: {result.stderr[-2000:]}"
            )
        for name in ("document.pdf", "spreadsheet.pdf"):
            path = root / name
            if not path.is_file():
                raise RuntimeError(
                    f"GDPVal judge Office conversion did not produce {name}"
                )
            with fitz.open(path) as pdf:
                if not pdf.page_count or "GDPVal" not in "".join(
                    p.get_text() for p in pdf
                ):
                    raise RuntimeError(f"GDPVal judge rendered an unreadable {name}")
