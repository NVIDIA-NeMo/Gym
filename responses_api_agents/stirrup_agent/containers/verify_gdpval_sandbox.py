# SPDX-FileCopyrightText: Copyright (c) 2026 NVIDIA CORPORATION & AFFILIATES. All rights reserved.
# SPDX-License-Identifier: Apache-2.0
"""Verify that a built GDPval sif really is the GDPval-AA v2 sandbox.

Runs INSIDE the container:

    apptainer exec --writable-tmpfs gdpval.sif \
        python /opt/gdpval/verify_gdpval_sandbox.py

A green ``gdpval.def`` build only proves the build steps exited 0. This checks
the finished image against the two vendored Artificial Analysis snapshots and
then actually exercises the toolchain, because the failure mode that costs a
whole eval run is a tool that is present but cannot produce a file.

Sections:
  pins      every one of the 419 published Python pins, at the exact version
  apt       the 762-entry published apt closure, split by whether the gap can
            affect behaviour or is an artefact of the base image differing
  binaries  every command line tool the task prompt advertises
  latex     style files from the curated TeX Live set resolve
  smoke     functional round-trips: each tool produces a non-empty artifact

Exit code is 0 only when no check in a fatal section failed.
"""

from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import sys
import tempfile
from importlib.metadata import PackageNotFoundError, version
from pathlib import Path


MANIFEST_DIR = Path("/opt/gdpval")
PY_MANIFEST = MANIFEST_DIR / "aa_v2_python_requirements.txt"
APT_MANIFEST = MANIFEST_DIR / "aa_v2_apt_closure.txt"

# Command line tools the GDPval task prompt tells the model it has. Advertising
# a tool the sandbox lacks is the expensive failure: the model plans around it
# and only discovers the gap mid-trajectory.
REQUIRED_BINARIES = [
    "python3",
    "pip",
    "git",
    "jq",
    "yq",
    "bash",
    "timeout",
    "libreoffice",
    "soffice",
    "pandoc",
    "tesseract",
    "pdftotext",
    "pdfimages",
    "gs",
    "ffmpeg",
    "ffprobe",
    "convert",
    "dot",
    "java",
    "gdalinfo",
    "unzip",
    "curl",
    "wget",
    "pdflatex",
    "xelatex",
    "lualatex",
    "latexmk",
    "biber",
    "kpsewhich",
]
# chromium is packaged under either name depending on the base image.
BINARY_ALTERNATIVES = {"chromium": ["chromium", "chromium-browser"]}

REQUIRED_STY = ["tikz.sty", "siunitx.sty", "pst-plot.sty", "amsmath.sty", "geometry.sty"]


class Report:
    def __init__(self) -> None:
        self.failures: list[str] = []
        self.warnings: list[str] = []

    def fail(self, section: str, msg: str) -> None:
        self.failures.append(f"[{section}] {msg}")
        print(f"  FAIL  {msg}", flush=True)

    def warn(self, section: str, msg: str) -> None:
        self.warnings.append(f"[{section}] {msg}")
        print(f"  warn  {msg}", flush=True)

    def ok(self, msg: str) -> None:
        print(f"  ok    {msg}", flush=True)


def header(name: str) -> None:
    print(f"\n=== {name} ===", flush=True)


def check_pins(rep: Report) -> None:
    header("pins: published Python manifest")
    if not PY_MANIFEST.exists():
        rep.fail("pins", f"{PY_MANIFEST} missing — image was not built from the aligned gdpval.def")
        return
    pins = [
        ln.strip()
        for ln in PY_MANIFEST.read_text(encoding="utf-8").splitlines()
        if ln.strip() and not ln.startswith("#")
    ]
    absent, wrong = [], []
    for pin in pins:
        name, want = pin.split("==", 1)
        try:
            got = version(name)
        except PackageNotFoundError:
            absent.append(name)
            continue
        if got != want:
            wrong.append(f"{name} want={want} got={got}")
    for name in sorted(absent):
        rep.fail("pins", f"not installed: {name}")
    for item in sorted(wrong):
        rep.fail("pins", f"version drift: {item}")
    if not absent and not wrong:
        rep.ok(f"all {len(pins)} published pins present at the published versions")


def check_apt(rep: Report, strict_apt: bool) -> None:
    header("apt: published system closure")
    if not APT_MANIFEST.exists():
        rep.fail("apt", f"{APT_MANIFEST} missing")
        return
    if not shutil.which("dpkg-query"):
        rep.fail("apt", "dpkg-query unavailable; cannot audit the system package set")
        return
    out = subprocess.run(
        ["dpkg-query", "-W", "-f=${Package}=${Version}\n"],
        capture_output=True,
        text=True,
        check=True,
    ).stdout
    installed = {}
    for line in out.splitlines():
        if "=" in line:
            pkg, ver = line.split("=", 1)
            installed[pkg] = ver

    wanted = {}
    for ln in APT_MANIFEST.read_text(encoding="utf-8").splitlines():
        ln = ln.strip()
        if not ln or ln.startswith("#") or "=" not in ln:
            continue
        pkg, ver = ln.split("=", 1)
        wanted[pkg] = ver

    missing = sorted(p for p in wanted if p not in installed)
    drift = sorted(
        f"{p} want={wanted[p]} got={installed[p]}" for p in wanted if p in installed and installed[p] != wanted[p]
    )

    # A missing library the toolchain never loads is noise; a missing tool is
    # not. Anything that ships a binary or a font is treated as behavioural.
    behavioural = [
        p
        for p in missing
        if not p.startswith("lib") or any(k in p for k in ("reoffice", "magick", "gdal", "proj", "geos"))
    ]
    incidental = [p for p in missing if p not in behavioural]

    print(f"  closure={len(wanted)} installed={len(installed)} missing={len(missing)} drift={len(drift)}", flush=True)
    for p in behavioural:
        if strict_apt:
            rep.fail("apt", f"missing (behavioural): {p}={wanted[p]}")
        else:
            rep.warn("apt", f"missing (behavioural): {p}={wanted[p]}")
    if incidental:
        rep.warn(
            "apt",
            f"missing transitive libs ({len(incidental)}): {', '.join(incidental[:25])}"
            + (" ..." if len(incidental) > 25 else ""),
        )
    if drift:
        rep.warn("apt", f"version drift on {len(drift)} packages, e.g. {'; '.join(drift[:8])}")
    if not missing and not drift:
        rep.ok("system package set matches the published closure exactly")


def check_binaries(rep: Report) -> None:
    header("binaries: tools the prompt advertises")
    for name in REQUIRED_BINARIES:
        if shutil.which(name):
            rep.ok(name)
        else:
            rep.fail("binaries", f"not on PATH: {name}")
    for label, alts in BINARY_ALTERNATIVES.items():
        found = next((a for a in alts if shutil.which(a)), None)
        if found:
            rep.ok(f"{label} (as {found})")
        else:
            rep.fail("binaries", f"not on PATH under any of {alts}: {label}")


def check_latex(rep: Report) -> None:
    header("latex: style files resolve")
    if not shutil.which("kpsewhich"):
        rep.fail("latex", "kpsewhich missing; cannot resolve style files")
        return
    for sty in REQUIRED_STY:
        res = subprocess.run(["kpsewhich", sty], capture_output=True, text=True)
        if res.returncode == 0 and res.stdout.strip():
            rep.ok(f"{sty} -> {res.stdout.strip()}")
        else:
            rep.fail("latex", f"unresolved style file: {sty}")


def _run(cmd: list[str], cwd: Path, timeout: int = 300) -> subprocess.CompletedProcess:
    return subprocess.run(cmd, cwd=cwd, capture_output=True, text=True, timeout=timeout)


def _nonempty(p: Path) -> bool:
    return p.exists() and p.stat().st_size > 0


def check_smoke(rep: Report, work: Path) -> None:
    header("smoke: each tool actually produces an artifact")

    def step(name: str):
        def deco(fn):
            try:
                fn()
            except Exception as exc:  # noqa: BLE001
                rep.fail("smoke", f"{name}: {type(exc).__name__}: {exc}")
            return fn

        return deco

    @step("matplotlib renders a PNG")
    def _mpl():
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        out = work / "plot.png"
        fig, ax = plt.subplots()
        ax.plot([0, 1, 2], [3, 1, 2])
        fig.savefig(out)
        plt.close(fig)
        assert _nonempty(out), "no PNG written"
        rep.ok("matplotlib renders a PNG")

    @step("python-docx writes a .docx")
    def _docx():
        import docx

        out = work / "doc.docx"
        d = docx.Document()
        d.add_heading("GDPval sandbox check", 0)
        d.add_paragraph("body text")
        d.save(out)
        assert _nonempty(out), "no docx written"
        rep.ok("python-docx writes a .docx")

    @step("openpyxl writes and reads an .xlsx")
    def _xlsx():
        import openpyxl

        out = work / "book.xlsx"
        wb = openpyxl.Workbook()
        wb.active["A1"] = 42
        wb.save(out)
        assert openpyxl.load_workbook(out).active["A1"].value == 42, "round-trip mismatch"
        rep.ok("openpyxl writes and reads an .xlsx")

    @step("python-pptx writes a .pptx")
    def _pptx():
        from pptx import Presentation

        out = work / "deck.pptx"
        prs = Presentation()
        prs.slides.add_slide(prs.slide_layouts[5]).shapes.title.text = "check"
        prs.save(out)
        assert _nonempty(out), "no pptx written"
        rep.ok("python-pptx writes a .pptx")

    @step("reportlab writes a PDF")
    def _reportlab():
        from reportlab.pdfgen import canvas

        out = work / "rl.pdf"
        c = canvas.Canvas(str(out))
        c.drawString(100, 700, "gdpval")
        c.save()
        assert _nonempty(out), "no PDF written"
        rep.ok("reportlab writes a PDF")

    @step("weasyprint renders HTML to PDF")
    def _weasy():
        from weasyprint import HTML

        out = work / "weasy.pdf"
        HTML(string="<h1>gdpval</h1><p>check</p>").write_pdf(str(out))
        assert _nonempty(out), "no PDF written"
        rep.ok("weasyprint renders HTML to PDF")

    @step("PyMuPDF reads a generated PDF")
    def _fitz():
        import fitz

        src = work / "rl.pdf"
        assert src.exists(), "reportlab step did not produce rl.pdf"
        with fitz.open(src) as doc:
            assert doc.page_count >= 1, "no pages"
            assert "gdpval" in doc[0].get_text(), "text not extracted"
        rep.ok("PyMuPDF reads a generated PDF")

    @step("pdfplumber extracts text")
    def _plumber():
        import pdfplumber

        with pdfplumber.open(work / "rl.pdf") as pdf:
            assert "gdpval" in (pdf.pages[0].extract_text() or ""), "text not extracted"
        rep.ok("pdfplumber extracts text")

    @step("libreoffice converts .docx to PDF")
    def _lo():
        out_dir = work / "lo"
        out_dir.mkdir(exist_ok=True)
        res = _run(
            [
                "soffice",
                "--headless",
                "--norestore",
                f"-env:UserInstallation=file://{work}/louser",
                "--convert-to",
                "pdf",
                "--outdir",
                str(out_dir),
                str(work / "doc.docx"),
            ],
            work,
            timeout=420,
        )
        pdf = out_dir / "doc.pdf"
        assert _nonempty(pdf), f"no PDF produced (rc={res.returncode}) {res.stdout[-400:]} {res.stderr[-400:]}"
        rep.ok("libreoffice converts .docx to PDF")

    @step("pandoc converts markdown to .docx")
    def _pandoc():
        src = work / "in.md"
        src.write_text("# Title\n\nSome **bold** text.\n", encoding="utf-8")
        out = work / "pandoc.docx"
        res = _run(["pandoc", str(src), "-o", str(out)], work)
        assert _nonempty(out), f"no docx produced (rc={res.returncode}) {res.stderr[-400:]}"
        rep.ok("pandoc converts markdown to .docx")

    @step("pdflatex compiles a document using tikz and siunitx")
    def _latex():
        tex = work / "t.tex"
        tex.write_text(
            r"\documentclass{article}\usepackage{tikz}\usepackage{siunitx}"
            r"\begin{document}\SI{3}{\kilo\gram}"
            r"\begin{tikzpicture}\draw (0,0)--(1,1);\end{tikzpicture}"
            r"\end{document}",
            encoding="utf-8",
        )
        res = _run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", "t.tex"], work, timeout=420)
        assert _nonempty(work / "t.pdf"), f"no PDF produced (rc={res.returncode}) {res.stdout[-600:]}"
        rep.ok("pdflatex compiles a document using tikz and siunitx")

    @step("ghostscript reprocesses a PDF")
    def _gs():
        out = work / "gs.pdf"
        res = _run(
            ["gs", "-q", "-dNOPAUSE", "-dBATCH", "-sDEVICE=pdfwrite", f"-sOutputFile={out}", str(work / "rl.pdf")],
            work,
        )
        assert _nonempty(out), f"no PDF produced (rc={res.returncode}) {res.stderr[-300:]}"
        rep.ok("ghostscript reprocesses a PDF")

    @step("pdftotext extracts text")
    def _pdftotext():
        out = work / "rl.txt"
        _run(["pdftotext", str(work / "rl.pdf"), str(out)], work)
        assert out.exists() and "gdpval" in out.read_text(errors="replace"), "text not extracted"
        rep.ok("pdftotext extracts text")

    @step("Pillow and opencv round-trip an image")
    def _img():
        import cv2
        import numpy as np
        from PIL import Image

        out = work / "img.png"
        Image.new("RGB", (64, 48), (10, 120, 200)).save(out)
        arr = cv2.imread(str(out))
        assert arr is not None and arr.shape == (48, 64, 3), f"opencv read back {None if arr is None else arr.shape}"
        assert np.asarray(Image.open(out)).shape == (48, 64, 3), "Pillow read back wrong shape"
        rep.ok("Pillow and opencv round-trip an image")

    @step("tesseract OCRs rendered text")
    def _ocr():
        import pytesseract
        from PIL import Image, ImageDraw

        img = Image.new("RGB", (420, 120), "white")
        ImageDraw.Draw(img).text((12, 40), "GDPVAL", fill="black")
        src = work / "ocr.png"
        img.save(src)
        # Upscale: the default bitmap font is too small for reliable OCR.
        big = img.resize((1680, 480), Image.LANCZOS)
        text = pytesseract.image_to_string(big).strip().upper()
        assert text, "tesseract returned no text at all"
        rep.ok(f"tesseract OCRs rendered text (read {text.splitlines()[0][:24]!r})")

    @step("ImageMagick converts PNG to JPEG")
    def _magick():
        out = work / "img.jpg"
        res = _run(["convert", str(work / "img.png"), str(out)], work)
        assert _nonempty(out), f"no JPEG produced (rc={res.returncode}) {res.stderr[-300:]}"
        rep.ok("ImageMagick converts PNG to JPEG")

    @step("ffmpeg encodes and probes a video")
    def _ffmpeg():
        out = work / "v.mp4"
        res = _run(
            [
                "ffmpeg",
                "-y",
                "-f",
                "lavfi",
                "-i",
                "testsrc=duration=1:size=128x96:rate=10",
                "-pix_fmt",
                "yuv420p",
                str(out),
            ],
            work,
        )
        assert _nonempty(out), f"no video produced (rc={res.returncode}) {res.stderr[-400:]}"
        probe = _run(
            ["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "default=nw=1:nk=1", str(out)], work
        )
        assert probe.returncode == 0 and probe.stdout.strip(), "ffprobe could not read the file"
        rep.ok("ffmpeg encodes and probes a video")

    @step("soundfile and librosa round-trip audio")
    def _audio():
        import librosa
        import numpy as np
        import soundfile as sf

        out = work / "a.wav"
        sr = 22050
        sf.write(out, np.sin(np.linspace(0, 220 * 2 * np.pi, sr)).astype("float32"), sr)
        y, got_sr = librosa.load(out, sr=None)
        assert got_sr == sr and y.size == sr, f"round-trip mismatch sr={got_sr} n={y.size}"
        rep.ok("soundfile and librosa round-trip audio")

    @step("graphviz renders a DOT graph")
    def _graphviz():
        import graphviz

        g = graphviz.Digraph()
        g.edge("a", "b")
        produced = Path(g.render(filename=str(work / "g"), format="png", cleanup=True))
        assert _nonempty(produced), "no PNG rendered"
        rep.ok("graphviz renders a DOT graph")

    @step("cairosvg rasterises SVG")
    def _cairosvg():
        import cairosvg

        out = work / "s.png"
        cairosvg.svg2png(
            bytestring=b'<svg xmlns="http://www.w3.org/2000/svg" width="40" height="40">'
            b'<rect width="40" height="40" fill="teal"/></svg>',
            write_to=str(out),
        )
        assert _nonempty(out), "no PNG written"
        rep.ok("cairosvg rasterises SVG")

    @step("pandas, polars and pyarrow exchange a frame")
    def _frames():
        import pandas as pd
        import polars as pl

        out = work / "f.parquet"
        pd.DataFrame({"a": [1, 2, 3], "b": ["x", "y", "z"]}).to_parquet(out)
        assert pl.read_parquet(out).shape == (3, 2), "polars read back the wrong shape"
        rep.ok("pandas, polars and pyarrow exchange a frame")

    @step("scikit-learn fits a model")
    def _sklearn():
        from sklearn.datasets import make_classification
        from sklearn.ensemble import RandomForestClassifier

        X, y = make_classification(n_samples=80, n_features=6, random_state=0)
        assert RandomForestClassifier(n_estimators=8, random_state=0).fit(X, y).score(X, y) > 0.5
        rep.ok("scikit-learn fits a model")

    @step("shapely and geopandas do a spatial op")
    def _geo():
        import geopandas as gpd
        from shapely.geometry import Point

        gdf = gpd.GeoDataFrame(geometry=[Point(0, 0).buffer(1)], crs="EPSG:4326")
        assert gdf.area.iloc[0] > 0, "degenerate geometry"
        rep.ok("shapely and geopandas do a spatial op")

    @step("gdalinfo reports on a raster")
    def _gdal():
        import numpy as np
        from PIL import Image

        tif = work / "r.tif"
        Image.fromarray((np.random.rand(32, 32) * 255).astype("uint8")).save(tif)
        res = _run(["gdalinfo", str(tif)], work)
        assert res.returncode == 0 and "Size is" in res.stdout, f"gdalinfo failed: {res.stderr[-300:]}"
        rep.ok("gdalinfo reports on a raster")

    @step("java runs (tabula backend)")
    def _java():
        res = _run(["java", "-version"], work)
        assert res.returncode == 0, f"java -version failed: {res.stderr[-200:]}"
        rep.ok("java runs (tabula backend)")

    @step("chromium starts headless")
    def _chromium():
        exe = shutil.which("chromium") or shutil.which("chromium-browser")
        assert exe, "chromium not installed"
        res = _run([exe, "--headless", "--no-sandbox", "--disable-gpu", "--version"], work, timeout=120)
        assert res.returncode == 0 and res.stdout.strip(), f"chromium did not report a version: {res.stderr[-300:]}"
        rep.ok(f"chromium starts headless ({res.stdout.strip()[:40]})")

    @step("sympy solves symbolically")
    def _sympy():
        import sympy

        x = sympy.Symbol("x")
        assert sympy.solve(x**2 - 4, x) == [-2, 2], "unexpected solution set"
        rep.ok("sympy solves symbolically")

    @step("magika identifies a file type")
    def _magika():
        from magika import Magika

        res = Magika().identify_path(work / "rl.pdf")
        assert res.output.label == "pdf", f"identified as {res.output.label}"
        rep.ok("magika identifies a file type")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument(
        "--strict-apt", action="store_true", help="treat behavioural apt gaps as failures rather than warnings"
    )
    ap.add_argument("--skip-smoke", action="store_true", help="manifest checks only")
    args = ap.parse_args()

    rep = Report()
    print(f"python: {sys.version.split()[0]}  platform: {sys.platform}  arch: {os.uname().machine}", flush=True)

    check_pins(rep)
    check_apt(rep, args.strict_apt)
    check_binaries(rep)
    check_latex(rep)
    if not args.skip_smoke:
        with tempfile.TemporaryDirectory(prefix="gdpval_verify_") as td:
            check_smoke(rep, Path(td))

    header("summary")
    print(f"  failures: {len(rep.failures)}   warnings: {len(rep.warnings)}", flush=True)
    for f in rep.failures:
        print(f"  FAIL {f}", flush=True)
    if rep.failures:
        print("\nRESULT: the sandbox does NOT match GDPval-AA v2", flush=True)
        return 1
    print("\nRESULT: sandbox verified against GDPval-AA v2", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
